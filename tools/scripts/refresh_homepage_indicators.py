#!/usr/bin/env python3
"""Refresh the annual NSO series used by the homepage indicator cards."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sqlite3
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
PUBLIC = ROOT / "data.mn/public/datasets"
CONTENT = ROOT / "data.mn/src/data/data"
CHARTS = ROOT / "data.mn/public/charts"
VERSIONS = ROOT / "tools/versions"
REGISTRY = ROOT / "tools/registry/data.db"


@dataclass(frozen=True)
class Indicator:
    dataset_id: str
    path: str
    source_updated_at: str
    start_year: int
    fixed_dimensions: int
    headers_en: tuple[str, ...]
    headers_mn: tuple[str, ...]
    fixed_values_en: tuple[str, ...] = ()
    fixed_values_mn: tuple[str, ...] = ()
    xlsx_names: tuple[str, str] | None = None


INDICATORS = (
    Indicator(
        "population-total-mongolia",
        "Population, household/1_Population, household/DT_NSO_0300_003V1.px",
        "2026-04-10T11:59:13",
        1956,
        2,
        ("year", "value"),
        ("он", "утга"),
    ),
    Indicator(
        "inflation-annual",
        "Economy, environment/Consumer Price Index/DT_NSO_0600_013V2.px",
        "2026-01-08T14:49:04",
        1991,
        1,
        ("year", "value"),
        ("он", "утга"),
    ),
    Indicator(
        "trade-total",
        "Economy, environment/Foreign Trade/DT_NSO_1400_001V1_year.px",
        "2026-04-13T09:30:49",
        1924,
        1,
        ("year", "value"),
        ("он", "утга"),
    ),
    Indicator(
        "unemployment-rate-national",
        "Labour, business/Decent work/DT_NSO_0400_049V1.px",
        "2026-06-03T16:22:42",
        2009,
        1,
        ("Category", "Year", "value"),
        ("Ангилал", "Он", "утга"),
        ("Total",),
        ("Бүгд",),
        (
            "unemployment-rate-national-en.xlsx",
            "unemployment-rate-national-mn.xlsx",
        ),
    ),
    Indicator(
        "labor-participation-national",
        (
            "Labour, business/Labour/"
            "LABOUR FORCE PARTICIPATION RATE, by sex, age group, aimags and the Capital/"
            "DT_NSO_0400_018V1_1.px"
        ),
        "2026-05-29T14:41:06",
        1992,
        3,
        ("year", "value"),
        ("он", "утга"),
    ),
    Indicator(
        "livestock-total",
        "Regional development/Livestock/DT_NSO_1001_109V1.px",
        "2026-07-21T17:37:54",
        1970,
        2,
        ("type_of_livestock", "region", "year", "value"),
        ("малын_төрөл", "бүс", "он", "утга"),
        ("Total", "Total"),
        ("Бүгд", "Улсын дүн"),
    ),
    Indicator(
        "salary-average-national",
        (
            "Labour, business/Wages/"
            "MONTHLY AVERAGE NOMINAL WAGES, by division of economic activities/"
            "DT_NSO_0400_022V1.px"
        ),
        "2026-05-04T16:27:24",
        2001,
        2,
        ("Year", "Salary (MNT 1000s)"),
        ("Он", "Цалин (МНТ 1000)"),
    ),
    Indicator(
        "household-income-total",
        (
            "Society, development/Household income and expenditure/"
            "DT_NSO_1900_001V1.px"
        ),
        "2026-05-29T12:44:23",
        1997,
        2,
        ("year", "value"),
        ("он", "утга"),
        xlsx_names=(
            "household-income-total-en.xlsx",
            "household-income-total-mn.xlsx",
        ),
    ),
)


def endpoint(language: str, path: str) -> str:
    encoded = "/".join(urllib.parse.quote(part, safe="") for part in path.split("/"))
    return f"https://data.1212.mn/api/v1/{language}/NSO/{encoded}"


def fetch_series(indicator: Indicator, language: str) -> list[tuple[int, float]]:
    url = endpoint(language, indicator.path)
    with urllib.request.urlopen(url, timeout=30) as response:
        metadata = json.load(response)

    variables = metadata["variables"]
    query = []
    for index, variable in enumerate(variables):
        values = variable["values"]
        selected = values[:1] if index < indicator.fixed_dimensions else values
        query.append(
            {
                "code": variable["code"],
                "selection": {"filter": "item", "values": selected},
            }
        )

    request = urllib.request.Request(
        url,
        data=json.dumps(
            {"query": query, "response": {"format": "json-stat2"}}
        ).encode(),
        headers={"Accept": "application/json", "Content-Type": "application/json"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        result = json.load(response)

    years = [int(year) for year in variables[-1]["valueTexts"]]
    values = result.get("value", [])
    if len(years) != len(values):
        raise ValueError(
            f"{indicator.dataset_id} returned {len(values)} values for {len(years)} years"
        )
    series = sorted(
        (year, value)
        for year, value in zip(years, values)
        if year >= indicator.start_year and value is not None
    )
    if not series or series[-1][0] != 2025:
        raise ValueError(f"{indicator.dataset_id} does not contain a 2025 value")
    return series


def output_rows(
    indicator: Indicator, language: str, series: list[tuple[int, float]]
) -> list[tuple[object, ...]]:
    fixed = indicator.fixed_values_en if language == "en" else indicator.fixed_values_mn
    return [(*fixed, year, value) for year, value in series]


def write_csv(path: Path, headers: tuple[str, ...], rows: list[tuple[object, ...]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)


def write_xlsx(path: Path, headers: tuple[str, ...], rows: list[tuple[object, ...]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for index, column in enumerate(sheet.columns, start=1):
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 40)
    workbook.save(path)


def write_bilingual_xlsx(
    path: Path,
    headers_en: tuple[str, ...],
    rows_en: list[tuple[object, ...]],
    headers_mn: tuple[str, ...],
    rows_mn: list[tuple[object, ...]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in (
        ("English", headers_en, rows_en),
        ("Mongolian", headers_mn, rows_mn),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for index, column in enumerate(sheet.columns, start=1):
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(width, 40)
    workbook.save(path)


def human_size(path: Path) -> str:
    size = path.stat().st_size
    if size < 1024:
        return f"{size} B"
    return f"{size / 1024:.1f} KB".replace(".0 KB", " KB")


def refresh_content(indicator: Indicator, version: int) -> None:
    for language in ("en", "mn"):
        mdx = CONTENT / language / f"{indicator.dataset_id}.mdx"
        text = mdx.read_text()
        text = re.sub(r"\b2024\b", "2025", text)
        text = re.sub(
            r"^dataVersion: \d+$", f"dataVersion: {version}", text, flags=re.MULTILINE
        )
        text = re.sub(
            r"^dataDate: .+$", "dataDate: 2025-12-31", text, flags=re.MULTILINE
        )
        csv_path = PUBLIC / f"{indicator.dataset_id}-{language}.csv"
        text = re.sub(
            rf'(?<=path: "/datasets/{re.escape(indicator.dataset_id)}-{language}\.csv"\n'
            rf'    format: "csv"\n    size: ")[^"]+',
            human_size(csv_path),
            text,
        )
        xlsx_pattern = re.compile(
            r'(?P<prefix>  - path: "/datasets/(?P<name>[^"]+\.xlsx)"\n'
            r'    format: "xlsx"\n    size: ")[^"]+'
        )

        def update_xlsx_size(match: re.Match[str]) -> str:
            path = PUBLIC / match.group("name")
            return match.group("prefix") + human_size(path)

        text = xlsx_pattern.sub(update_xlsx_size, text)
        mdx.write_text(text)

        chart = CHARTS / f"{indicator.dataset_id}-{language}.json"
        chart.write_text(re.sub(r"\b2024\b", "2025", chart.read_text()))


def snapshot(
    indicator: Indicator,
    version: int,
    csv_en: Path,
    csv_mn: Path,
    xlsx_files: list[Path],
) -> Path:
    destination = VERSIONS / indicator.dataset_id / f"v{version}"
    destination.mkdir(parents=True, exist_ok=True)
    for source in (csv_en, csv_mn, *xlsx_files):
        shutil.copy2(source, destination / source.name)
    metadata = {
        "dataset_id": indicator.dataset_id,
        "version": version,
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "data_as_of": "2025",
        "source_updated_at": indicator.source_updated_at,
        "files": [csv_en.name, csv_mn.name, *(path.name for path in xlsx_files)],
    }
    (destination / "metadata.json").write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False) + "\n"
    )
    return destination


def update_registry(
    indicator: Indicator,
    version: int,
    destination: Path,
    csv_en: Path,
    row_count: int,
) -> None:
    fetched_at = datetime.now(timezone.utc).isoformat()
    digest = hashlib.sha256(csv_en.read_bytes()).hexdigest()
    with sqlite3.connect(REGISTRY) as connection:
        connection.execute(
            """UPDATE datasets
               SET status='active', current_version=?, data_as_of='2025',
                   source_ref=?, source_path=?, source_updated_at=?, last_checked_at=?,
                   last_fetched_at=?
               WHERE id=?""",
            (
                version,
                indicator.path.rsplit("/", 1)[-1],
                indicator.path,
                indicator.source_updated_at,
                fetched_at,
                fetched_at,
                indicator.dataset_id,
            ),
        )
        connection.execute(
            "DELETE FROM versions WHERE dataset_id=? AND version=?",
            (indicator.dataset_id, version),
        )
        connection.execute(
            """INSERT INTO versions
               (dataset_id, version, data_hash, data_path, row_count, column_count,
                change_type, change_summary, source_updated_at, source_raw_files,
                fetched_at)
               VALUES (?, ?, ?, ?, ?, ?, 'source_update', ?, ?, ?, ?)""",
            (
                indicator.dataset_id,
                version,
                digest,
                str(destination.relative_to(ROOT)),
                row_count,
                len(indicator.headers_en),
                "Refreshed official NSO annual series through 2025",
                indicator.source_updated_at,
                json.dumps(
                    [
                        f"{indicator.dataset_id}-en.csv",
                        f"{indicator.dataset_id}-mn.csv",
                    ]
                ),
                fetched_at,
            ),
        )
        connection.execute(
            """INSERT INTO activity_log
               (action, dataset_id, source_id, status, message, triggered_by)
               VALUES ('update', ?, 'nso-1212', 'success',
                       'Refreshed annual homepage indicator through 2025', 'manual')""",
            (indicator.dataset_id,),
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--dataset",
        action="append",
        choices=[indicator.dataset_id for indicator in INDICATORS],
        help="Refresh only this dataset (repeatable); defaults to all homepage indicators",
    )
    args = parser.parse_args()
    selected = set(args.dataset or ())

    with sqlite3.connect(REGISTRY) as connection:
        versions = {
            dataset_id: max(current_version or 0, recorded_version or 0)
            for dataset_id, current_version, recorded_version in connection.execute(
                """SELECT d.id, d.current_version, MAX(v.version)
                   FROM datasets d
                   LEFT JOIN versions v ON v.dataset_id=d.id
                   GROUP BY d.id, d.current_version"""
            )
        }

    for indicator in INDICATORS:
        if selected and indicator.dataset_id not in selected:
            continue
        series_en = fetch_series(indicator, "en")
        series_mn = fetch_series(indicator, "mn")
        if series_en != series_mn:
            raise ValueError(f"Bilingual values differ for {indicator.dataset_id}")

        rows_en = output_rows(indicator, "en", series_en)
        rows_mn = output_rows(indicator, "mn", series_mn)
        csv_en = PUBLIC / f"{indicator.dataset_id}-en.csv"
        csv_mn = PUBLIC / f"{indicator.dataset_id}-mn.csv"
        write_csv(csv_en, indicator.headers_en, rows_en)
        write_csv(csv_mn, indicator.headers_mn, rows_mn)
        if indicator.xlsx_names:
            xlsx_files = [PUBLIC / name for name in indicator.xlsx_names]
            write_xlsx(xlsx_files[0], indicator.headers_en, rows_en)
            write_xlsx(xlsx_files[1], indicator.headers_mn, rows_mn)
            generic_xlsx = PUBLIC / f"{indicator.dataset_id}.xlsx"
            if generic_xlsx.exists():
                write_bilingual_xlsx(
                    generic_xlsx,
                    indicator.headers_en,
                    rows_en,
                    indicator.headers_mn,
                    rows_mn,
                )
                xlsx_files.append(generic_xlsx)
        else:
            xlsx_files = [PUBLIC / f"{indicator.dataset_id}.xlsx"]
            write_xlsx(xlsx_files[0], indicator.headers_en, rows_en)
            for language, headers, rows in (
                ("en", indicator.headers_en, rows_en),
                ("mn", indicator.headers_mn, rows_mn),
            ):
                language_xlsx = PUBLIC / f"{indicator.dataset_id}-{language}.xlsx"
                if language_xlsx.exists():
                    write_xlsx(language_xlsx, headers, rows)
                    xlsx_files.append(language_xlsx)

        version = int(versions[indicator.dataset_id] or 0) + 1
        refresh_content(indicator, version)
        destination = snapshot(indicator, version, csv_en, csv_mn, xlsx_files)
        update_registry(indicator, version, destination, csv_en, len(rows_en))
        print(
            f"{indicator.dataset_id}: {series_en[-1][0]}={series_en[-1][1]} "
            f"(version {version})"
        )


if __name__ == "__main__":
    main()

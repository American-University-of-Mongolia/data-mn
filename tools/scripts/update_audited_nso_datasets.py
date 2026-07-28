#!/usr/bin/env python3
"""Regenerate the NSO datasets identified by the July 2026 freshness audit.

The script consumes bilingual raw CSVs produced by datamn-source-nso and
updates the published derivatives, version snapshots, MDX freshness metadata,
and registry records in one repeatable operation.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
PUBLIC = ROOT / "data.mn/public/datasets"
CONTENT = ROOT / "data.mn/src/data/data"
CHARTS = ROOT / "data.mn/public/charts"
VERSIONS = ROOT / "tools/versions"
REGISTRY = ROOT / "tools/registry/data.db"
TODAY = "2026-07-27"

SOURCE_UPDATED = {
    "nso-gdp-by-economic-activity": "2026-04-20T16:32:03",
    "nso-bop-monthly": "2026-07-08T15:25:08",
    "nso-temperature-by-station": "2026-07-22T16:12:00",
    "industrial-production-national": "2026-06-22T11:51:13",
}

GDP_SPLITS = [
    "gdp-nominal",
    "gdp-real",
    "gdp-usd",
    "gdp-growth-rate",
    "gdp-by-sector",
    "gdp-sector-trends",
]
BOP_SPLITS = [
    "bop-current-account",
    "bop-trade-balance",
    "bop-services-balance",
    "bop-fdi-net",
    "bop-reserve-assets",
    "bop-remittances",
]
TEMPERATURE_SPLITS = [
    "temperature-ulaanbaatar",
    "temperature-regional",
    "temperature-extremes-ulaanbaatar",
    "temperature-anomaly",
]


def clean_strings(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    for column in frame.select_dtypes(include=["object"]).columns:
        frame[column] = frame[column].astype(str).str.strip()
    return frame


def write_csv_pair(dataset_id: str, en: pd.DataFrame, mn: pd.DataFrame) -> None:
    en.to_csv(PUBLIC / f"{dataset_id}-en.csv", index=False)
    mn.to_csv(PUBLIC / f"{dataset_id}-mn.csv", index=False)
    write_xlsx(PUBLIC / f"{dataset_id}.xlsx", en)


def write_xlsx(path: Path, frame: pd.DataFrame) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Data")
        sheet = writer.sheets["Data"]
        for index, column in enumerate(frame.columns, start=1):
            values = frame[column].fillna("").astype(str)
            length = max(len(str(column)), int(values.str.len().max())) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(length, 40)


def aligned_language_frames(en_path: Path, mn_path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    en = clean_strings(pd.read_csv(en_path))
    mn = clean_strings(pd.read_csv(mn_path))
    if len(en) != len(mn):
        raise ValueError(f"Bilingual row mismatch: {en_path} ({len(en)}) vs {mn_path} ({len(mn)})")
    return en, mn


def regenerate_gdp(raw_root: Path) -> None:
    en, mn = aligned_language_frames(
        raw_root / "gdp/nso-0500-001v1-en.csv",
        raw_root / "gdp/nso-0500-001v1-mn.csv",
    )
    en.columns = ["indicator", "division", "year", "value"]
    mn.columns = ["indicator", "division", "year", "value"]

    def simple(indicator: str, dataset_id: str) -> None:
        mask = (en.indicator == indicator) & (en.division == "Total") & en.value.notna()
        out_en = en.loc[mask, ["year", "value"]].sort_values("year")
        out_mn = mn.loc[mask, ["year", "value"]].sort_values("year")
        out_mn.columns = ["он", "утга"]
        write_csv_pair(dataset_id, out_en, out_mn)

    simple("GDP, at current prices", "gdp-nominal")
    simple("GDP, at 2015 constant prices", "gdp-real")
    simple("GDP, at current price, thousand US dollars", "gdp-usd")
    simple("Annual changes, by percent", "gdp-growth-rate")

    current = en.indicator.eq("GDP, at current prices")
    latest_year = int(en.loc[current & en.value.notna(), "year"].astype(int).max())
    sector_mask = current & en.year.astype(int).eq(latest_year) & en.division.ne("Total") & en.value.notna()
    sector_en = en.loc[sector_mask, ["division", "value"]].rename(columns={"division": "sector"})
    sector_mn = mn.loc[sector_mask, ["division", "value"]].rename(
        columns={"division": "салбар", "value": "утга"}
    )
    write_csv_pair("gdp-by-sector", sector_en, sector_mn)

    sectors = {
        "Agriculture, forestry, fishing and hunting": "Agriculture",
        "Mining and quarrying": "Mining",
        "Manufacturing": "Manufacturing",
        "Construction": "Construction",
        "Wholesale and retail trade; repair of motor vehicles and motorcycles": "Trade",
        "Transportation and storage": "Transport",
    }
    trends_mask = current & en.division.isin(sectors) & en.value.notna()
    trends_en = en.loc[trends_mask, ["division", "year", "value"]].copy()
    trends_en["division"] = trends_en["division"].map(sectors)
    trends_en.columns = ["sector", "year", "value"]
    trends_en = trends_en.sort_values(["sector", "year"])
    trends_mn = mn.loc[trends_mask, ["division", "year", "value"]].copy()
    trends_mn.columns = ["салбар", "он", "утга"]
    trends_mn = trends_mn.sort_values(["салбар", "он"])
    write_csv_pair("gdp-sector-trends", trends_en, trends_mn)

    snapshot("nso-gdp-by-economic-activity", 2, en_path=raw_root / "gdp/nso-0500-001v1-en.csv",
             mn_path=raw_root / "gdp/nso-0500-001v1-mn.csv")


def regenerate_bop(raw_root: Path) -> None:
    en, mn = aligned_language_frames(
        raw_root / "bop/nso-0100-001v10-en.csv",
        raw_root / "bop/nso-0100-001v10-mn.csv",
    )
    en.columns = ["indicator", "month", "value"]
    mn.columns = ["indicator", "month", "value"]

    def simple(indicator: str, dataset_id: str) -> None:
        mask = en.indicator.eq(indicator) & en.value.notna()
        out_en = en.loc[mask, ["month", "value"]].sort_values("month")
        out_mn = mn.loc[mask, ["month", "value"]].sort_values("month")
        out_mn.columns = ["сар", "утга"]
        write_csv_pair(dataset_id, out_en, out_mn)

    simple("I. CURRENT ACCOUNT", "bop-current-account")
    simple("2. Services", "bop-services-balance")
    simple("1. Direct investment", "bop-fdi-net")
    simple("V. RESERVE ASSETS", "bop-reserve-assets")
    simple("of which: Personal transfers", "bop-remittances")

    exports = en.loc[en.indicator.eq("1.1 Export FOB (credit)"), ["month", "value"]]
    imports = en.loc[en.indicator.eq("1.2 Import FOB (debit)"), ["month", "value"]]
    trade = exports.merge(imports, on="month", suffixes=("_exports", "_imports")).dropna()
    trade.columns = ["month", "exports", "imports"]
    trade["net"] = trade.exports - trade.imports
    trade = trade.sort_values("month")
    trade_mn = trade.copy()
    trade_mn.columns = ["сар", "экспорт", "импорт", "цэвэр"]
    write_csv_pair("bop-trade-balance", trade, trade_mn)

    snapshot("nso-bop-monthly", 2, en_path=raw_root / "bop/nso-0100-001v10-en.csv",
             mn_path=raw_root / "bop/nso-0100-001v10-mn.csv")


def normalize_month(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, format="%Y-%m").dt.strftime("%Y-%m")


def regenerate_temperature(raw_root: Path) -> None:
    en, mn = aligned_language_frames(
        raw_root / "temperature/nso-2400-022v2-en.csv",
        raw_root / "temperature/nso-2400-022v2-mn.csv",
    )
    en.columns = ["indicator", "station", "month", "value"]
    mn.columns = ["indicator", "station", "month", "value"]
    en["month"] = normalize_month(en.month)
    mn["month"] = normalize_month(mn.month)

    average = en.indicator.eq("Average air temperature")
    ub = en.station.eq("Ulaanbaatar")
    mask = average & ub & en.value.notna()
    out_en = en.loc[mask, ["month", "value"]].rename(columns={"value": "temperature"}).sort_values("month")
    out_mn = mn.loc[mask, ["month", "value"]].rename(
        columns={"month": "сар", "value": "температур"}
    ).sort_values("сар")
    write_csv_pair("temperature-ulaanbaatar", out_en, out_mn)

    stations = ["Ulaanbaatar", "Darkhan", "Dalanzadgad", "Choibalsan", "Khovd"]
    mask = average & en.station.isin(stations) & en.value.notna()
    out_en = en.loc[mask, ["month", "station", "value"]].rename(
        columns={"value": "temperature"}
    ).sort_values(["month", "station"])
    out_mn = mn.loc[mask, ["month", "station", "value"]].rename(
        columns={"month": "сар", "station": "станц", "value": "температур"}
    ).sort_values(["сар", "станц"])
    write_csv_pair("temperature-regional", out_en, out_mn)

    labels_en = {
        "Average air temperature": "Average",
        "Maximum temperature": "Maximum",
        "Minimum temperature": "Minimum",
    }
    labels_mn = {"Average": "Дундаж", "Maximum": "Хамгийн их", "Minimum": "Хамгийн бага"}
    mask = en.indicator.isin(labels_en) & ub & en.value.notna()
    out_en = en.loc[mask, ["month", "indicator", "value"]].copy()
    out_en["indicator"] = out_en.indicator.map(labels_en)
    out_en.columns = ["month", "indicator", "temperature"]
    out_en = out_en.sort_values(["month", "indicator"])
    out_mn = out_en.copy()
    out_mn["indicator"] = out_mn.indicator.map(labels_mn)
    out_mn.columns = ["сар", "үзүүлэлт", "температур"]
    write_csv_pair("temperature-extremes-ulaanbaatar", out_en, out_mn)

    anomaly = en.indicator.eq("Comparison with multi-year (1981-2010)")
    mask = anomaly & ub & en.value.notna()
    out_en = en.loc[mask, ["month", "value"]].rename(columns={"value": "anomaly"}).sort_values("month")
    out_mn = mn.loc[mask, ["month", "value"]].rename(
        columns={"month": "сар", "value": "хазайлт"}
    ).sort_values("сар")
    write_csv_pair("temperature-anomaly", out_en, out_mn)

    snapshot("nso-temperature-by-station", 1,
             en_path=raw_root / "temperature/nso-2400-022v2-en.csv",
             mn_path=raw_root / "temperature/nso-2400-022v2-mn.csv")


def regenerate_industrial(raw_root: Path) -> None:
    en, mn = aligned_language_frames(
        raw_root / "industrial/nso-1100-013v1-en.csv",
        raw_root / "industrial/nso-1100-013v1-mn.csv",
    )
    en.columns = ["commodity", "year", "value"]
    mn.columns = ["commodity", "year", "value"]
    selected = {
        "Coal (thous.t)": ("Coal", "Нүүрс", "thousand tonnes", "мян.тонн"),
        "Copper concentrate with 35% /thous.t/": (
            "Copper concentrate (35%)", "Зэсийн баяжмал (35%)", "thousand tonnes", "мян.тонн"
        ),
        "Gold /kg/": ("Gold", "Алт", "kg", "кг"),
        "Electricity /mln.k.W.h/": ("Electricity", "Цахилгаан эрчим хүч", "million kWh", "сая кВт.ц"),
        "Crude oil /thous.barrel/": ("Crude oil", "Газрын тос", "thousand barrels", "мян.баррель"),
        "Iron ore /thous.t/": ("Iron ore", "Төмрийн хүдэр", "thousand tonnes", "мян.тонн"),
        "Molybdenium concentrate with 47 % /t/": (
            "Molybdenum concentrate (47%)", "Молибдений баяжмал (47%)", "tonnes", "тонн"
        ),
    }
    mask = en.commodity.isin(selected) & en.value.notna()
    rows_en = []
    rows_mn = []
    for index in en.index[mask]:
        name_en, name_mn, unit_en, unit_mn = selected[en.at[index, "commodity"]]
        rows_en.append((int(en.at[index, "year"]), name_en, unit_en, en.at[index, "value"]))
        rows_mn.append((int(mn.at[index, "year"]), name_mn, unit_mn, mn.at[index, "value"]))
    out_en = pd.DataFrame(rows_en, columns=["year", "commodity", "unit", "value"]).sort_values(
        ["year", "commodity"]
    )
    out_mn = pd.DataFrame(
        rows_mn, columns=["он", "бараа_бүтээгдэхүүн", "хэмжих_нэгж", "утга"]
    ).sort_values(["он", "бараа_бүтээгдэхүүн"])
    write_csv_pair("industrial-production-national", out_en, out_mn)
    snapshot("industrial-production-national", 2,
             en_path=raw_root / "industrial/nso-1100-013v1-en.csv",
             mn_path=raw_root / "industrial/nso-1100-013v1-mn.csv")


def snapshot(dataset_id: str, version: int, *, en_path: Path, mn_path: Path) -> None:
    destination = VERSIONS / dataset_id / f"v{version}"
    destination.mkdir(parents=True, exist_ok=True)
    shutil.copy2(en_path, destination / f"{dataset_id}-en.csv")
    shutil.copy2(mn_path, destination / f"{dataset_id}-mn.csv")
    metadata = {
        "dataset_id": dataset_id,
        "version": version,
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "source_updated_at": SOURCE_UPDATED[dataset_id],
        "source_files": [f"{dataset_id}-en.csv", f"{dataset_id}-mn.csv"],
    }
    (destination / "version.json").write_text(json.dumps(metadata, indent=2) + "\n")


def snapshot_published(dataset_ids: list[str], version: int, source_updated_at: str) -> None:
    for dataset_id in dataset_ids:
        destination = VERSIONS / dataset_id / f"v{version}"
        destination.mkdir(parents=True, exist_ok=True)
        files = [
            PUBLIC / f"{dataset_id}-en.csv",
            PUBLIC / f"{dataset_id}-mn.csv",
            PUBLIC / f"{dataset_id}.xlsx",
        ]
        for source in files:
            shutil.copy2(source, destination / source.name)
        metadata = {
            "dataset_id": dataset_id,
            "version": version,
            "fetched_at": datetime.now(timezone.utc).isoformat(),
            "source_updated_at": source_updated_at,
            "files": [source.name for source in files],
        }
        (destination / "metadata.json").write_text(json.dumps(metadata, indent=2) + "\n")


def refresh_mdx_and_charts(dataset_ids: list[str], latest: str, version: int) -> None:
    latest_year = latest[:4]
    for dataset_id in dataset_ids:
        for language in ("en", "mn"):
            path = CONTENT / language / f"{dataset_id}.mdx"
            text = path.read_text()
            text = re.sub(r"^dataVersion: \d+$", f"dataVersion: {version}", text, flags=re.MULTILINE)
            text = re.sub(r"^dataDate: .+$", f"dataDate: {latest}", text, flags=re.MULTILINE)
            lines = []
            for line in text.splitlines():
                if line.startswith(("title:", "excerpt:", "  title=")):
                    line = re.sub(r"\b(?:2024|2025)\b", latest_year, line)
                lines.append(line)
            text = "\n".join(lines) + ("\n" if text.endswith("\n") else "")
            path.write_text(text)
            chart = CHARTS / f"{dataset_id}-{language}.json"
            if chart.exists():
                chart_text = chart.read_text()
                chart_text = re.sub(r"(?<=-)2024(?=\b)|(?<=-)2025(?=\b)", latest_year, chart_text)
                chart.write_text(chart_text)


def create_missing_english_xlsx() -> None:
    dataset_ids = [
        "meat-production-historical-by-type",
        "infant-mortality-rate-national",
        "meat-production-historical-total",
        "hospital-beds-national",
        "meat-production-by-region",
    ]
    for dataset_id in dataset_ids:
        source = PUBLIC / f"{dataset_id}-en.csv"
        destination = PUBLIC / f"{dataset_id}-en.xlsx"
        write_xlsx(destination, pd.read_csv(source))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def update_registry() -> None:
    fetched = datetime.now(timezone.utc).isoformat()
    groups = [
        ("nso-gdp-by-economic-activity", GDP_SPLITS, "2025", 2),
        ("nso-bop-monthly", BOP_SPLITS, "2026-05", 2),
        ("nso-temperature-by-station", TEMPERATURE_SPLITS, "2026-06", 1),
        ("industrial-production-national", [], "2025", 2),
    ]
    with sqlite3.connect(REGISTRY) as connection:
        for parent, splits, data_as_of, version in groups:
            ids = [parent, *splits]
            placeholders = ",".join("?" for _ in ids)
            connection.execute(
                f"""UPDATE datasets
                    SET status='active', current_version=?, data_as_of=?,
                        source_updated_at=?, last_checked_at=?, last_fetched_at=?
                    WHERE id IN ({placeholders})""",
                (version, data_as_of, SOURCE_UPDATED[parent], fetched, fetched, *ids),
            )
            data_path = str(VERSIONS / parent / f"v{version}")
            raw_en = Path(data_path) / f"{parent}-en.csv"
            connection.execute(
                "DELETE FROM versions WHERE dataset_id=? AND version=?",
                (parent, version),
            )
            connection.execute(
                """INSERT INTO versions
                   (dataset_id, version, data_hash, data_path, row_count, column_count,
                    change_type, change_summary, source_updated_at, source_raw_files)
                   VALUES (?, ?, ?, ?, ?, ?, 'source_update', ?, ?, ?)""",
                (
                    parent,
                    version,
                    sha256(raw_en),
                    str(Path(data_path).relative_to(ROOT)),
                    len(pd.read_csv(raw_en)),
                    len(pd.read_csv(raw_en, nrows=0).columns),
                    f"Refreshed source data through {data_as_of}",
                    SOURCE_UPDATED[parent],
                    json.dumps([raw_en.name, f"{parent}-mn.csv"]),
                ),
            )
            connection.execute(
                """INSERT INTO activity_log
                   (action, dataset_id, source_id, status, message, triggered_by)
                   VALUES ('update', ?, 'nso-1212', 'success', ?, 'manual')""",
                (parent, f"Updated source data and {len(splits)} derived datasets through {data_as_of}"),
            )

        # The two nested table IDs remain valid; only their API paths were incomplete.
        connection.execute(
            """UPDATE datasets SET source_path=?
               WHERE id='labor-participation-national'""",
            ("Labour, business/Labour/LABOUR FORCE PARTICIPATION RATE, by sex, age group, aimags and the Capital/DT_NSO_0400_018V1_1.px",),
        )
        connection.execute(
            """UPDATE datasets SET source_path=?
               WHERE id='salary-average-national'""",
            ("Labour, business/Wages/MONTHLY AVERAGE NOMINAL WAGES, by division of economic activities/DT_NSO_0400_022V1.px",),
        )
        # V1 was retired; V3 preserves the type/region/year dimensions.
        connection.execute(
            """UPDATE datasets SET source_ref=?, source_path=?, source_updated_at=?
               WHERE id='hospital-beds-national'""",
            (
                "DT_NSO_2100_005V3.px",
                "Education, health/Main indicators for Health sector/DT_NSO_2100_005V3.px",
                "2026-05-06T15:41:29",
            ),
        )

        source_metadata = {
            "household-income-by-location": (
                "DT_NSO_1900_001V1.px", "Society, development/Household income and expenditure",
                "2026-05-29T12:44:23"
            ),
            "household-income-total": (
                "DT_NSO_1900_001V1.px", "Society, development/Household income and expenditure",
                "2026-05-29T12:44:23"
            ),
            "nso-deaths-by-region-monthly": (
                "DT_NSO_2100_027V2.px", "Education, health/Births, deaths",
                "2026-07-09T18:30:38"
            ),
            "nso-deaths-monthly": (
                "DT_NSO_2100_027V2.px", "Education, health/Births, deaths",
                "2026-07-09T18:30:38"
            ),
            "nso-exports-by-country-top": (
                "DT_NSO_1400_006V3.px", "Economy, environment/Foreign Trade",
                "2026-04-14T17:27:45"
            ),
            "nso-foreign-trade-annual": (
                "DT_NSO_1400_001V1_year.px", "Economy, environment/Foreign Trade",
                "2026-04-13T09:30:49"
            ),
            "nso-foreign-trade-monthly": (
                "DT_NSO_1400_003V1.px", "Economy, environment/Foreign Trade",
                "2026-07-09T12:25:28"
            ),
            "nso-imports-by-country-top": (
                "DT_NSO_1400_010V3.px", "Economy, environment/Foreign Trade",
                "2026-04-15T11:02:21"
            ),
            "nso-live-births-by-region-monthly": (
                "DT_NSO_2100_018V5.px", "Education, health/Births, deaths",
                "2026-07-09T18:29:03"
            ),
            "nso-live-births-monthly": (
                "DT_NSO_2100_018V5.px", "Education, health/Births, deaths",
                "2026-07-09T18:29:03"
            ),
            "nso-marriages-divorces-annual": (
                "DT_NSO_0300_018V2.px", "Population, household/2_Regular movement of population",
                "2026-05-05T11:17:03"
            ),
            "nso-population-age-sex": (
                "DT_NSO_0300_003V1.px", "Population, household/1_Population, household",
                "2026-05-01T14:52:02"
            ),
            "trade-exports-imports": (
                "DT_NSO_1400_001V1_year.px", "Economy, environment/Foreign Trade",
                "2026-04-13T09:30:49"
            ),
            "trade-total": (
                "DT_NSO_1400_001V1_year.px", "Economy, environment/Foreign Trade",
                "2026-04-13T09:30:49"
            ),
        }
        for dataset_id, (source_ref, source_path, updated) in source_metadata.items():
            connection.execute(
                """UPDATE datasets
                   SET source_ref=?, source_path=?, source_updated_at=?, last_checked_at=?
                   WHERE id=?""",
                (source_ref, source_path, updated, fetched, dataset_id),
            )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--raw-root",
        type=Path,
        required=True,
        help="Directory containing gdp/, bop/, temperature/, and industrial/ raw fetches",
    )
    args = parser.parse_args()
    regenerate_gdp(args.raw_root)
    regenerate_bop(args.raw_root)
    regenerate_temperature(args.raw_root)
    regenerate_industrial(args.raw_root)
    refresh_mdx_and_charts(GDP_SPLITS, "2025-12-31", 2)
    refresh_mdx_and_charts(BOP_SPLITS, "2026-05-31", 2)
    refresh_mdx_and_charts(TEMPERATURE_SPLITS, "2026-06-30", 1)
    refresh_mdx_and_charts(["industrial-production-national"], "2025-12-31", 2)
    snapshot_published(GDP_SPLITS, 2, SOURCE_UPDATED["nso-gdp-by-economic-activity"])
    snapshot_published(BOP_SPLITS, 2, SOURCE_UPDATED["nso-bop-monthly"])
    snapshot_published(TEMPERATURE_SPLITS, 1, SOURCE_UPDATED["nso-temperature-by-station"])
    create_missing_english_xlsx()
    update_registry()


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""
Rebuild the Ulaanbaatar month-over-month CPI splits from the parent snapshot.

Parent: nso-cpi-ulaanbaatar-mom (NSO DT_NSO_0600_003V4.px), snapshot in
tools/versions/nso-cpi-ulaanbaatar-mom/v1. NSO publishes the same MoM
change on two reference bases whose baskets differ, so the series are
joined with a marked break instead of being silently spliced:

  2020=100  for months up to and including BREAK_LAST_2020 (published
            history is unchanged)
  2023=100  for every later month

The break is marked on the chart (dashed rule + label at the first
2023=100 month, see patch_chart), in the XLSX (cell comment on that month's
row) and in the page excerpt. The CSVs stay standard time/[category]/value
files: the download standard has no annotation column. The 2020=100 series
was discontinued after 2025-01; where both bases exist (2023-02..2025-01)
they differ by up to ~0.5 pp.

Splits written (EN + MN):
  cpi-monthly-ulaanbaatar     Overall index          (+ -all copy)
  cpi-food-ulaanbaatar        Food and non-alcoholic beverages
  cpi-by-category-ulaanbaatar mean of grouped COICOP divisions:
      Food & Dining       alcohol/tobacco, food, restaurants/hotels
      Housing & Utilities communication, furnishings, housing
      Transport           transport
      Services            education, health, insurance, recreation
      Other Goods         clothing, miscellaneous

Usage:
    python3 tools/scripts/build_cpi_ulaanbaatar.py           # write CSVs + charts
    python3 tools/scripts/build_cpi_ulaanbaatar.py --check   # exit 1 if files differ

"""

import argparse
import json
import re
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

ROOT = Path(__file__).resolve().parents[2]
SNAPSHOT = ROOT / "tools/versions/nso-cpi-ulaanbaatar-mom/v1"
PUBLIC = ROOT / "data.mn/public"
BREAK_LAST_2020 = "2025-01"
DATASETS = ["cpi-monthly-ulaanbaatar", "cpi-food-ulaanbaatar", "cpi-by-category-ulaanbaatar"]
BASES = {"old": "2020=100", "new": "2023=100"}

CATEGORIES = {  # EN name -> (MN name, source groups)
    "Food & Dining": ("Хүнс, хоол", ["Alcoholic beverages and tobacco", "Food and non-alcaholic beverages",
                                      "Restaurants and hotels"]),
    "Housing & Utilities": ("Орон сууц, хэрэгсэл", ["Communication", "Furnishings, household equipment and tools",
                                                      "Housing, water, electricity and fuels"]),
    "Other Goods": ("Бусад бараа", ["Clothing, footwear and cloth", "Miscellaneous goods and services"]),
    "Services": ("Үйлчилгээ", ["Education services", "Health, medical care and services",
                                "Insurance and financial services", "Recreation and culture"]),
    "Transport": ("Тээвэр", ["Transport"]),
}
BREAK_LABEL = {"en": "Basis change: 2020=100 → 2023=100", "mn": "Суурь өөрчлөгдсөн: 2020=100 → 2023=100"}
# chart label: two short lines, right-aligned to the rule (the break is near the end of the series)
BREAK_LINES = {"en": ["Basis change", "2020=100 → 2023=100"], "mn": ["Суурь өөрчлөгдсөн", "2020=100 → 2023=100"]}


def joined_wide() -> pd.DataFrame:
    """Month x group MoM changes: 2020=100 up to the break, 2023=100 after."""
    d = pd.read_csv(SNAPSHOT / "nso-0600-003v4-en.csv").dropna(subset=["value"])
    d["Group"] = d.Group.str.strip()
    out = []
    for key, base in BASES.items():
        w = d[d["Reference year"] == base].pivot(index="Month", columns="Group", values="value")
        w = w[w.index <= BREAK_LAST_2020] if key == "old" else w[w.index > BREAK_LAST_2020]
        w["basis"] = base  # used for the break month only; not written to CSVs
        out.append(w)
    wide = pd.concat(out).sort_index()
    wide.index.name = "month"
    return wide


def first_new_month(wide: pd.DataFrame) -> str:
    return wide.index[wide.basis == BASES["new"]][0]


def frames(wide: pd.DataFrame):
    """{dataset_id: (en_df, mn_df)}."""
    month = wide.index
    def one(series):
        en = pd.DataFrame({"month": month, "value": series.round(3).values})
        mn = en.rename(columns={"month": "сар", "value": "утга"})
        return en, mn
    result = {
        "cpi-monthly-ulaanbaatar": one(wide["Overall index"]),
        "cpi-food-ulaanbaatar": one(wide["Food and non-alcaholic beverages"]),
    }
    rows = []
    for en_name in sorted(CATEGORIES):
        mn_name, groups = CATEGORIES[en_name]
        values = wide[groups].mean(axis=1).round(3)
        rows += [(m, en_name, mn_name, v) for m, v in zip(month, values)]
    cat = pd.DataFrame(rows, columns=["month", "category", "ангилал", "value"])
    result["cpi-by-category-ulaanbaatar"] = (
        cat[["month", "category", "value"]],
        cat[["month", "ангилал", "value"]].rename(columns={"month": "сар", "value": "утга"}),
    )
    return result


def is_marker(layer: dict) -> bool:
    """The basis-change marker layers carry inline data; data layers don't."""
    return "values" in layer.get("data", {})


def patch_chart(spec: dict, first_new: str, lang: str, csv_name: str, version: int) -> dict:
    """Add the basis-change marker; idempotent.

    Marker layers carry their own data and override every channel the
    top-level encoding would otherwise leak into them (y, color).
    """
    spec["data"]["url"] = f"/datasets/{csv_name}?v={version}"  # CDN caches CSVs 4h
    spec["layer"] = [layer for layer in spec["layer"] if not is_marker(layer)]
    grey = {"value": "#64748b"}
    x = {"field": "month", "type": "temporal"}
    marker = {"data": {"values": [{"month": first_new, "label": BREAK_LABEL[lang]}]}}
    spec["layer"] += [
        {**marker,
         "mark": {"type": "rule", "strokeDash": [5, 4], "strokeWidth": 1.5},
         "encoding": {"x": x, "y": {"value": 0}, "y2": {"value": "height"}, "color": grey}},
        {**marker,
         "mark": {"type": "text", "align": "right", "dx": -6, "dy": -4, "baseline": "bottom", "fontSize": 12},
         "encoding": {"x": x, "y": {"value": "height"}, "color": grey, "text": {"value": BREAK_LINES[lang]}}},
    ]
    return spec


def render(dataset_id: str, en: pd.DataFrame, mn: pd.DataFrame, first_new: str):
    """{path: text} for one dataset."""
    files = {}
    csv = lambda df: df.to_csv(index=False)
    files[PUBLIC / f"datasets/{dataset_id}-en.csv"] = csv(en)
    files[PUBLIC / f"datasets/{dataset_id}-mn.csv"] = csv(mn)
    if dataset_id == "cpi-monthly-ulaanbaatar":
        files[PUBLIC / f"datasets/{dataset_id}-all-en.csv"] = csv(en)
        files[PUBLIC / f"datasets/{dataset_id}-all-mn.csv"] = csv(mn)
    version = 3 if dataset_id == "cpi-by-category-ulaanbaatar" else 2
    for lang in ("en", "mn"):
        path = PUBLIC / f"charts/{dataset_id}-{lang}.json"
        spec = patch_chart(json.loads(path.read_text()), first_new, lang, f"{dataset_id}-{lang}.csv", version)
        files[path] = json.dumps(spec, indent=2, ensure_ascii=False) + "\n"
    return files


def write_xlsx(dataset_id, en, mn, first_new, lang_labels=BREAK_LABEL):
    """Bilingual wide XLSX (same layout as rebuild_downloads) + a comment on the break row."""
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from rebuild_downloads import pivot_long, write_bilingual_xlsx
    path = PUBLIC / f"datasets/{dataset_id}.xlsx"
    sheets = []
    for df in (en, mn):
        cols = list(df.columns)
        sheets.append(pivot_long(df, cols[0], cols[1], cols[2]) if len(cols) == 3 else df.copy())
    write_bilingual_xlsx(path, *sheets)
    wb = load_workbook(path)
    for ws, lang in zip(wb.worksheets, ("en", "mn")):
        for row in ws.iter_rows(min_row=2, max_col=1):
            if row[0].value == first_new:
                row[0].comment = Comment(lang_labels[lang], "data.mn")
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--check", action="store_true", help="report files that differ; write nothing")
    args = parser.parse_args()
    wide = joined_wide()
    first_new = first_new_month(wide)
    changed = 0
    for dataset_id, (en, mn) in frames(wide).items():
        if not args.check:
            write_xlsx(dataset_id, en, mn, first_new)
        for path, text in render(dataset_id, en, mn, first_new).items():
            if not path.exists() or path.read_text() != text:
                changed += 1
                print(("DIFFERS  " if args.check else "WROTE    ") + str(path.relative_to(ROOT)))
                if not args.check:
                    path.write_text(text)
    if not args.check:  # keep title/description year spans in step with the data
        from validate_title_years import CONTENT, check_page
        for dataset_id in DATASETS:
            for lang in ("en", "mn"):
                check_page(CONTENT / lang / f"{dataset_id}.mdx", fix=True)
    print(f"{wide.index[0]} .. {wide.index[-1]}, break at {first_new}; {changed} file(s) "
          f"{'differ' if args.check else 'written'}")
    return 1 if args.check and changed else 0


if __name__ == "__main__":
    sys.exit(main())
